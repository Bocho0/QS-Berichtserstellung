"""
Gemeinsamer Kern der QS-Berichtserstellung - wird sowohl vom lokalen
CLI-Skript (build_report.py) als auch von der Vercel-Serverfunktion
(api/generate.py) verwendet, damit es nur EINE gepflegte Quelle für die
Formatierungslogik gibt.

WICHTIG - neue Architektur (Muster-Umsetzung):
Alle STATISCHEN Formatierungs-/Struktur-Änderungen aus dem Muster-Handoff
(Schriften, Farben, Ränder, Linien statt Rahmen, Spaltenbreiten, verbundene
Kopfzellen, Wetterband-Container, Ampel-Legende, Deckblatt-Kopfzeilen-Fix)
sind bereits FEST in `api/vorlage.docx` eingebaut (siehe
`scripts/build_template.py`, mit dem die Vorlage aus dem Original erzeugt
wurde). Dieses Modul hier befüllt die Vorlage nur noch mit Werten -
Datenmodell und Feldnamen sind gegenüber der Vorgängerversion unverändert.
"""
import os, copy, base64, re, math
import docx
from docx.shared import Pt, Emu, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement
from lxml import etree
from xml.sax.saxutils import escape as xml_escape
from PIL import Image

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
EMU_PER_IN = 914400

COL_TEXT = '1C1C1C'
COL_LABEL_DOC = '8A8A8A'
COL_MUTED_ID = '5A5A5A'
COL_LABEL_COVER = '7A7A7A'
COL_TABLE_HEAD = '6A6A6A'
COL_LINE_DARK = '000000'
COL_LINE_LIGHT = 'DEDEDE'
COL_FOOTER = '8A8A8A'

# Fotogröße (Muster-Handoff Abschnitt 3).
MAX_BOX_IN = 2.48
MAX_BOX_WIDTH_IN = 3.31
LINE_HEIGHT_PT = 9 * 1.2
LABEL_TAB_DXA = 851  # 15mm - Feststellungs-Raster (muss zur Vorlage passen)


def w(tag):
    return f'{{{W}}}{tag}'


# ---------------------------------------------------------------------------
# Kleine Helfer
# ---------------------------------------------------------------------------

def decode_photo(data_url, out_path):
    header, b64 = data_url.split(',', 1)
    with open(out_path, 'wb') as f:
        f.write(base64.b64decode(b64))
    return out_path


def fit_box(w_px, h_px, max_h_in=MAX_BOX_IN, max_w_in=MAX_BOX_WIDTH_IN):
    ratio = w_px / h_px
    ratio = max(0.65, min(1.7, ratio))
    h_in = max_h_in
    w_in = max_h_in * ratio
    if w_in > max_w_in:
        w_in = max_w_in
        h_in = max_w_in / ratio
    return Emu(int(w_in * EMU_PER_IN)), Emu(int(h_in * EMU_PER_IN))


def format_datum(raw):
    raw = (raw or '').strip()
    m = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})$', raw)
    if m:
        d, mo, y = m.groups()
        return f"{int(d):02d}.{int(mo):02d}.{y}"
    return raw


def format_temp(raw):
    raw = (raw or '').strip()
    if not raw:
        return raw
    if '°' in raw or raw.lower().endswith('c'):
        return raw
    return raw + '°C'


WEEKDAYS_DE = ['Mo', 'Di', 'Mi', 'Do', 'Fr', 'Sa', 'So']


def format_datum_mit_wochentag(raw):
    raw = (raw or '').strip()
    m = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})$', raw)
    if not m:
        return raw
    d, mo, y = (int(x) for x in m.groups())
    try:
        import datetime
        wd = WEEKDAYS_DE[datetime.date(y, mo, d).weekday()]
        return f"{wd}, {d:02d}.{mo:02d}.{y}"
    except Exception:
        return f"{d:02d}.{mo:02d}.{y}"


def set_run_font(run, base_run, color=None):
    if base_run is not None:
        run.font.name = base_run.font.name or 'Barlow'
        run.font.size = base_run.font.size or Pt(9)
    else:
        run.font.name = 'Barlow'
        run.font.size = Pt(9)
    if color:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)
    elif base_run is not None and base_run.font.color and base_run.font.color.type is not None and base_run.font.color.rgb:
        run.font.color.rgb = base_run.font.color.rgb


def clear_cell(cell):
    tc = cell._tc
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    if len(cell.paragraphs) == 0:
        tc.append(parse_xml(f'<w:p xmlns:w="{W}"/>'))


# ---------------------------------------------------------------------------
# Deckblatt: Werte in die bereits fertig formatierten "Label\tWert"-Absätze
# einsetzen (Formatierung kommt komplett aus der Vorlage).
# ---------------------------------------------------------------------------

def set_cover_value(doc, label, value):
    for p in doc.paragraphs:
        if p.runs and p.runs[0].text.strip() == label:
            if len(p.runs) > 1:
                p.runs[1].text = '\t' + value
            else:
                r = p.add_run('\t' + value)
                set_run_font(r, p.runs[0], color=COL_TEXT)
            return True
    return False


def set_seitenanzahl_field(doc, literal_total=None):
    """Setzt statt einer festen Zahl ein echtes Word-Feld (NUMPAGES) ein,
    damit die Seitenanzahl auch nach späteren Bearbeitungen in Word
    automatisch stimmt. Bei literal_total (zweiteiliger Export) wird
    stattdessen die vorab berechnete Gesamtseitenzahl beider Teile fest
    eingetragen."""
    for p in doc.paragraphs:
        if p.runs and p.runs[0].text.strip() == 'Seitenanzahl':
            value_run = p.runs[1] if len(p.runs) > 1 else p.add_run()
            ref_rPr = value_run._r.find(qn('w:rPr'))
            rpr_xml = (etree.tostring(ref_rPr, encoding='unicode') if ref_rPr is not None
                       else f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/></w:rPr>')
            value_run.text = '\t'
            def make_run(inner):
                return parse_xml(f'<w:r xmlns:w="{W}">{rpr_xml}{inner}</w:r>')
            anchor = value_run._r
            if literal_total:
                for rn in [make_run(f'<w:t>{int(literal_total)}</w:t>'),
                           make_run('<w:t xml:space="preserve"> Seiten</w:t>')]:
                    anchor.addnext(rn); anchor = rn
                return
            for rn in [
                make_run('<w:fldChar w:fldCharType="begin" w:dirty="true"/>'),
                make_run('<w:instrText>NUMPAGES   \\* MERGEFORMAT</w:instrText>'),
                make_run('<w:fldChar w:fldCharType="separate"/>'),
                make_run('<w:t>1</w:t>'),
                make_run('<w:fldChar w:fldCharType="end"/>'),
                make_run('<w:t xml:space="preserve"> Seiten</w:t>'),
            ]:
                anchor.addnext(rn); anchor = rn
            return


# ---------------------------------------------------------------------------
# Wetterband: 3 Platzhalter-Runs aus der Vorlage per Text ersetzen.
# ---------------------------------------------------------------------------

def fill_wetterband(doc, bk):
    datum_txt = format_datum_mit_wochentag(bk.get('datum', ''))
    temp_min = format_temp(bk.get('tempMin'))
    temp_max = format_temp(bk.get('tempMax'))
    min_txt = f'Min. {temp_min}' if temp_min else ''
    max_txt = f'Max. {temp_max}' if temp_max else ''
    replacements = {
        'WETTERBAND_DATUM': datum_txt,
        'WETTERBAND_MIN': min_txt,
        'WETTERBAND_MAX': max_txt,
    }
    for p in doc.paragraphs:
        for r in p.runs:
            for placeholder, value in replacements.items():
                if placeholder in r.text:
                    r.text = r.text.replace(placeholder, value)


# ---------------------------------------------------------------------------
# Verteiler-Tabelle (Formatierung bereits in der Vorlage, hier nur Text)
# ---------------------------------------------------------------------------

def fill_verteiler(vt, verteiler_raw):
    rows_text = [r.strip() for r in (verteiler_raw or '').split('\n') if r.strip()]
    for i, row_text in enumerate(rows_text):
        if i + 1 >= len(vt.rows):
            break
        parts = [p.strip() for p in row_text.split(';')]
        parts += [''] * (4 - len(parts))
        row = vt.rows[i + 1]
        cells = row.cells
        # cells[1] und cells[2] zeigen wegen gridSpan=2 (Firma) auf dieselbe
        # Zelle - Python-docx dedupliziert das nicht automatisch, daher
        # explizit auf eindeutige Zellen abbilden.
        unique_cells = []
        seen = set()
        for c in cells:
            if id(c._tc) not in seen:
                unique_cells.append(c)
                seen.add(id(c._tc))
        for cell, val in zip(unique_cells, parts[:4]):
            base_run = None
            for p in cell.paragraphs:
                if p.runs:
                    base_run = p.runs[0]
                    break
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r.text = ''
            target_p = cell.paragraphs[0]
            r = target_p.add_run(val)
            set_run_font(r, base_run, color=COL_TEXT)
    used = min(len(rows_text), len(vt.rows) - 1)
    for row in list(vt.rows[used + 1:]):
        vt._tbl.remove(row._tr)


# ---------------------------------------------------------------------------
# Rahmentermine-Tabelle: EINE Muster-Datenzeile aus der Vorlage pro Eintrag
# klonen (wie die Dokumentationstabelle).
# ---------------------------------------------------------------------------

def fill_rahmentermine(rt, bk):
    BAUTENSTAND_LABELS = {
        'nicht-begonnen': 'Nicht begonnen',
        'in-bearbeitung': 'In Bearbeitung',
        'fertiggestellt': 'Fertiggestellt',
    }
    RAHMEN_COLOR = {
        'im-termin': 'B2CB7F',
        'plus2': 'F8A764',
        'plus4': 'F95649',
    }
    rows_text = [r.strip() for r in (bk.get('rahmentermine') or '').split('\n') if r.strip()]
    if not rows_text:
        # Muster-Zeile einfach entfernen, wenn keine Rahmentermine vorliegen.
        for row in list(rt.rows)[1:]:
            rt._tbl.remove(row._tr)
        return

    template_tr = copy.deepcopy(rt.rows[1]._tr)
    base_run = None
    for p in rt.rows[1].cells[1].paragraphs:
        if p.runs:
            base_run = p.runs[0]
            break
    for row in list(rt.rows)[1:]:
        rt._tbl.remove(row._tr)

    for row_text in rows_text:
        parts = [p.strip() for p in row_text.split(';')]
        parts += [''] * (8 - len(parts))
        haus, gewerk, tp_von, tp_bis, bautenstand_code, prozent, status, prognose = parts[:8]
        bautenstand_display = BAUTENSTAND_LABELS.get(bautenstand_code, bautenstand_code)
        prozent_display = f'{prozent}%' if prozent else ''
        col_text = {0: haus, 1: gewerk, 2: tp_von, 3: tp_bis, 4: bautenstand_display,
                    5: prozent_display, 7: prognose}
        fill = RAHMEN_COLOR.get(status.lower(), 'FFFFFF')

        new_tr = copy.deepcopy(template_tr)
        rt._tbl.append(new_tr)
        row = rt.rows[-1]
        for col_idx, val in col_text.items():
            if col_idx >= len(row.cells):
                continue
            cell = row.cells[col_idx]
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r.text = ''
            target_p = cell.paragraphs[0]
            r = target_p.add_run(val)
            color = COL_MUTED_ID if col_idx == 0 else COL_TEXT
            set_run_font(r, base_run, color=color)
            if col_idx == 5:
                target_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if len(row.cells) > 6:
            color_cell = row.cells[6]
            tcPr = color_cell._tc.get_or_add_tcPr()
            old_shd = tcPr.find(qn('w:shd'))
            if old_shd is not None:
                tcPr.remove(old_shd)
            shd = tcPr.makeelement(qn('w:shd'), {})
            shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
            tcPr.append(shd)


# ---------------------------------------------------------------------------
# Dokumentations-/Veranlassungs-Zellen: Prototyp-Absätze aus der Vorlage
# (PROTO_LABEL/PROTO_VALUE/PROTO_CONT/PROTO_STAND) je Feld duplizieren.
# ---------------------------------------------------------------------------

def _find_proto_paragraphs(cell):
    line_p = cont_p = stand_p = None
    for p in cell.paragraphs:
        t = p.text
        if 'PROTO_LABEL' in t:
            line_p = p
        elif 'PROTO_CONT' in t:
            cont_p = p
        elif 'PROTO_STAND' in t:
            stand_p = p
    return line_p, cont_p, stand_p


def estimate_wrapped_extra_lines(text, col_width_dxa=4590 - LABEL_TAB_DXA, font_pt=9):
    if not text:
        return 0
    avail_pt = max(20, col_width_dxa / 20)
    chars_per_line = max(10, int(avail_pt / (font_pt * 0.5)))
    return max(0, math.ceil(len(text) / chars_per_line) - 1)


def build_text_cell(cell, entry, base_run_unused, image_height_pt):
    proto_line, proto_cont, proto_stand = _find_proto_paragraphs(cell)
    tc = cell._tc

    def clone_line(label, value, first=False):
        new_p = copy.deepcopy(proto_line._p)
        for r_el in new_p.findall(qn('w:r')):
            t_el = r_el.find(qn('w:t'))
            if t_el is not None and t_el.text and 'PROTO_LABEL' in t_el.text:
                t_el.text = label
            elif t_el is not None and t_el.text and 'PROTO_VALUE' in t_el.text:
                t_el.text = value
        if not first:
            # Nur die allererste Zeile (Gewerk) bekommt den grossen
            # Vorabstand zur Zellenoberkante - alle weiteren Feldzeilen
            # sollen eng aufeinander folgen (Muster: row-gap ~0,6mm).
            pPr = new_p.find(qn('w:pPr'))
            sp = pPr.find(qn('w:spacing')) if pPr is not None else None
            if sp is not None:
                sp.set(qn('w:before'), '12')
        return new_p

    def clone_cont(text):
        new_p = copy.deepcopy(proto_cont._p)
        for r_el in new_p.findall(qn('w:r')):
            t_el = r_el.find(qn('w:t'))
            if t_el is not None:
                t_el.text = text
        return new_p

    def clone_stand(text):
        new_p = copy.deepcopy(proto_stand._p)
        for r_el in new_p.findall(qn('w:r')):
            t_el = r_el.find(qn('w:t'))
            if t_el is not None:
                t_el.text = text
        return new_p

    new_paragraphs = []
    line_count = 0
    is_first_field = True

    def add_field(label, value):
        nonlocal line_count, is_first_field
        new_paragraphs.append(clone_line(label, value, first=is_first_field))
        is_first_field = False
        line_count += 1 + estimate_wrapped_extra_lines(value)

    add_field('Gewerk', entry.get('gewerk', ''))
    add_field('Ort', entry.get('ort', ''))
    if entry.get('bauteil'):
        add_field('Bauteil', entry['bauteil'])

    info = entry.get('arbeiten', '') or ''
    info_lines = [l for l in info.split('\n')] if info.strip() else []
    if info_lines:
        add_field('Info', info_lines[0])
        for extra in info_lines[1:]:
            new_paragraphs.append(clone_cont(extra))
            line_count += 1 + estimate_wrapped_extra_lines(extra)

    if entry.get('material'):
        add_field('Material', entry['material'])
    if entry.get('type') == 'veranlassung' and entry.get('verantwortlich'):
        add_field('Verantwortlichkeit', entry['verantwortlich'])

    FINE_TUNE_PT = 6
    used_pt = 16 + line_count * LINE_HEIGHT_PT
    space_before_pt = max(14, image_height_pt + 16 - used_pt - LINE_HEIGHT_PT - FINE_TUNE_PT)
    stand_p = clone_stand(f"Stand: {format_datum(entry.get('datum', ''))}")
    stand_pPr = stand_p.find(qn('w:pPr'))
    if stand_pPr is None:
        stand_pPr = stand_p.makeelement(qn('w:pPr'), {})
        stand_p.insert(0, stand_pPr)
    old_sp = stand_pPr.find(qn('w:spacing'))
    if old_sp is not None:
        stand_pPr.remove(old_sp)
    sp_el = stand_pPr.makeelement(qn('w:spacing'), {})
    sp_el.set(qn('w:before'), str(int(space_before_pt * 20)))
    stand_pPr.append(sp_el)

    # Erste Zeile bekommt den ueblichen 16dxa-Vorabstand der Vorlage
    # (schon im Prototyp enthalten) - Original-Prototypen jetzt entfernen
    # und durch die generierten Absaetze ersetzen.
    anchor = proto_line._p
    for new_p in new_paragraphs:
        anchor.addprevious(new_p)
    anchor.addprevious(stand_p)
    proto_line._p.getparent().remove(proto_line._p)
    proto_cont._p.getparent().remove(proto_cont._p)
    proto_stand._p.getparent().remove(proto_stand._p)


def build_image_cell(cell, photo_path):
    clear_cell(cell)
    tcPr = cell._tc.get_or_add_tcPr()
    old_valign = tcPr.find(qn('w:vAlign'))
    if old_valign is not None:
        tcPr.remove(old_valign)
    valign = tcPr.makeelement(qn('w:vAlign'), {})
    valign.set(qn('w:val'), 'center')
    tcPr.append(valign)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run()
    with Image.open(photo_path) as im:
        w_px, h_px = im.size
    w_emu, h_emu = fit_box(w_px, h_px)
    r.add_picture(photo_path, width=w_emu, height=h_emu)
    return h_emu / EMU_PER_IN * 72


def set_nr_cell(cell, nr):
    clear_cell(cell)
    tcPr = cell._tc.get_or_add_tcPr()
    old_valign = tcPr.find(qn('w:vAlign'))
    if old_valign is not None:
        tcPr.remove(old_valign)
    valign = tcPr.makeelement(qn('w:vAlign'), {})
    valign.set(qn('w:val'), 'top')
    tcPr.append(valign)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run(str(nr))
    from docx.shared import RGBColor
    r.font.name = 'Barlow'
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(COL_MUTED_ID)


# ---------------------------------------------------------------------------
# Hauptfunktion
# ---------------------------------------------------------------------------

def find_table_by_header(doc, header_text):
    """Robuste Tabellenzuordnung über den Text der ersten Kopfzellezeile
    statt über eine feste Position in doc.tables - unempfindlich gegenüber
    zusätzlichen (rein dekorativen) Tabellen wie der Ampel-Legende."""
    for t in doc.tables:
        if t.rows and t.rows[0].cells:
            cell_text = t.rows[0].cells[0].text.strip()
            if cell_text.startswith(header_text):
                return t
    raise ValueError(f"Tabelle mit Kopfzeile '{header_text}' nicht gefunden.")


def build(template_path, data, out_path, tmp_dir='/tmp/report_photos', only_content=False):
    """data: bereits geparstes dict (nicht Dateipfad!).
    only_content=True: Titelblatt, Verteiler, Wetter, Rahmentermine und
    Inhaltsverzeichnis werden am Ende komplett entfernt - übrig bleiben nur
    die Feststellungen/Veranlassungen (Teil 2 eines zweiteiligen Berichts).
    """
    bk = data.get('berichtskopf', {})
    entries = data.get('entries', [])
    verlauf = data.get('verlauf', [])
    start_nr = data.get('startNr') or 1
    try:
        start_nr = int(start_nr)
    except (TypeError, ValueError):
        start_nr = 1

    doc = docx.Document(template_path)

    verteiler_table = find_table_by_header(doc, 'Name')
    rahmentermine_table = find_table_by_header(doc, 'Haus')
    dokumentation_table = find_table_by_header(doc, 'Nr.')

    # ------------------------------------------------------------------
    # Inhaltsverzeichnis + Bookmarks (dynamisch, da abhängig von der
    # tatsächlichen Seitenaufteilung beim Öffnen in Word)
    # ------------------------------------------------------------------
    def set_outline_level(paragraph, level):
        pPr = paragraph._p.get_or_add_pPr()
        old = pPr.find(qn('w:outlineLvl'))
        if old is not None:
            pPr.remove(old)
        el = pPr.makeelement(qn('w:outlineLvl'), {})
        el.set(qn('w:val'), str(level))
        pPr.append(el)

    anchor_p = pagebreak_p = dokumentation_p = rahmentermine_p = None
    content_pagebreak_p = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('Anlagen'):
            anchor_p = p
        if t in ('Rahmentermine', 'Dokumentation'):
            set_outline_level(p, 0)
            if t == 'Dokumentation':
                dokumentation_p = p
            if t == 'Rahmentermine':
                rahmentermine_p = p
        if pagebreak_p is None:
            for br in p._p.findall('.//' + qn('w:br')):
                if br.get(qn('w:type')) == 'page':
                    pagebreak_p = p
                    break
        if rahmentermine_p is not None and dokumentation_p is None and content_pagebreak_p is None:
            for br in p._p.findall('.//' + qn('w:br')):
                if br.get(qn('w:type')) == 'page':
                    content_pagebreak_p = p
                    break

    def add_bookmark(paragraph, name, bm_id):
        p_el = paragraph._p
        pPr = p_el.find(qn('w:pPr'))
        start = parse_xml(f'<w:bookmarkStart xmlns:w="{W}" w:id="{bm_id}" w:name="{name}"/>')
        end = parse_xml(f'<w:bookmarkEnd xmlns:w="{W}" w:id="{bm_id}"/>')
        if pPr is not None:
            pPr.addnext(end)
            pPr.addnext(start)
        else:
            p_el.insert(0, end)
            p_el.insert(0, start)

    if rahmentermine_p is not None:
        add_bookmark(rahmentermine_p, 'bm_rahmentermine', 901)
    if dokumentation_p is not None:
        add_bookmark(dokumentation_p, 'bm_dokumentation', 902)

    if anchor_p is not None and pagebreak_p is not None:
        toc_heading_rpr_xml = (f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/>'
                                f'<w:caps w:val="1"/><w:spacing w:val="24"/><w:sz w:val="15"/><w:szCs w:val="15"/>'
                                f'<w:color w:val="{COL_TABLE_HEAD}"/></w:rPr>')
        toc_text_rpr_xml = f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
        toc_num_rpr_xml = f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="18"/><w:szCs w:val="18"/><w:color w:val="{COL_LABEL_COVER}"/></w:rPr>'

        toc_heading = parse_xml(f'''<w:p xmlns:w="{W}">
          <w:pPr><w:spacing w:before="1600" w:after="100"/>
            <w:pBdr><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{COL_LINE_DARK}"/></w:pBdr></w:pPr>
          <w:r>{toc_heading_rpr_xml}<w:t>Inhaltsverzeichnis</w:t></w:r>
        </w:p>''')
        pagebreak_p._p.addprevious(toc_heading)

        sec = doc.sections[0]
        avail_dxa = sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips

        def make_toc_line(label, bookmark_name):
            return parse_xml(f'''<w:p xmlns:w="{W}">
              <w:pPr>
                <w:tabs><w:tab w:val="right" w:leader="none" w:pos="{avail_dxa}"/></w:tabs>
                <w:spacing w:before="102" w:after="102"/>
                <w:pBdr><w:bottom w:val="single" w:sz="3" w:space="2" w:color="{COL_LINE_LIGHT}"/></w:pBdr>
              </w:pPr>
              <w:r>{toc_text_rpr_xml}<w:t xml:space="preserve">{xml_escape(label)}</w:t></w:r>
              <w:r>{toc_text_rpr_xml}<w:tab/></w:r>
              <w:r>{toc_num_rpr_xml}<w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>
              <w:r>{toc_num_rpr_xml}<w:instrText xml:space="preserve"> PAGEREF {bookmark_name} \\h </w:instrText></w:r>
              <w:r>{toc_num_rpr_xml}<w:fldChar w:fldCharType="separate"/></w:r>
              <w:r>{toc_num_rpr_xml}<w:t>2</w:t></w:r>
              <w:r>{toc_num_rpr_xml}<w:fldChar w:fldCharType="end"/></w:r>
            </w:p>''')

        pagebreak_p._p.addprevious(make_toc_line('Rahmentermine', 'bm_rahmentermine'))
        pagebreak_p._p.addprevious(make_toc_line('Dokumentation', 'bm_dokumentation'))

        toc_line_y = sec.page_height.twips - sec.bottom_margin.twips
        toc_bottom_spacer = parse_xml(f'''<w:p xmlns:w="{W}">
          <w:pPr>
            <w:framePr w:w="{avail_dxa}" w:h="20" w:hRule="atLeast" w:hAnchor="margin" w:vAnchor="page" w:x="0" w:y="{toc_line_y}" w:wrap="none"/>
            <w:pBdr><w:top w:val="single" w:sz="6" w:space="1" w:color="{COL_LINE_DARK}"/></w:pBdr>
          </w:pPr>
        </w:p>''')
        pagebreak_p._p.addprevious(toc_bottom_spacer)

        first_page_footer = doc.sections[0].first_page_footer
        for fp in first_page_footer.paragraphs:
            fpPr = fp._p.find(qn('w:pPr'))
            if fpPr is not None:
                old_bdr = fpPr.find(qn('w:pBdr'))
                if old_bdr is not None:
                    fpPr.remove(old_bdr)

    # ------------------------------------------------------------------
    # Verlauf vorheriger Berichte (dynamisch, auf der Rahmentermine-Seite)
    # ------------------------------------------------------------------
    verlauf_anchor = content_pagebreak_p if content_pagebreak_p is not None else dokumentation_p
    if verlauf and verlauf_anchor is not None:
        vref_rpr_xml = f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:b/><w:sz w:val="21"/><w:szCs w:val="21"/><w:color w:val="{COL_TEXT}"/></w:rPr>'
        verlauf_heading = parse_xml(f'''<w:p xmlns:w="{W}">
          <w:pPr><w:spacing w:before="681" w:after="170"/></w:pPr>
          <w:r>{vref_rpr_xml}<w:t>Verlauf vorheriger Berichte</w:t></w:r>
        </w:p>''')
        verlauf_anchor._p.addprevious(verlauf_heading)

        sec = doc.sections[0]
        avail_dxa = sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips
        col_nr, col_zeit = 1300, 2600
        col_zus = avail_dxa - col_nr - col_zeit

        def vcell(text, width, header=False, muted=False):
            if header:
                rpr = f'<w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="15"/><w:szCs w:val="15"/><w:spacing w:val="4"/><w:color w:val="{COL_TABLE_HEAD}"/>'
                border = f'<w:tcBorders><w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COL_LINE_DARK}"/></w:tcBorders>'
            else:
                color = COL_MUTED_ID if muted else COL_TEXT
                rpr = f'<w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="18"/><w:szCs w:val="18"/><w:color w:val="{color}"/>'
                border = f'<w:tcBorders><w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/><w:bottom w:val="single" w:sz="3" w:space="0" w:color="{COL_LINE_LIGHT}"/></w:tcBorders>'
            return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{border}</w:tcPr>'
                    f'<w:p><w:pPr><w:spacing w:before="40" w:after="40"/></w:pPr><w:r><w:rPr>{rpr}</w:rPr>'
                    f'<w:t xml:space="preserve">{xml_escape(text)}</w:t></w:r></w:p></w:tc>')

        rows_xml = '<w:tr>' + vcell('Nr.', col_nr, True) + vcell('Zeitraum', col_zeit, True) + vcell('Zusammenfassung', col_zus, True) + '</w:tr>'
        for v in verlauf:
            von, bis = v.get('von', ''), v.get('bis', '')
            nr_text = str(von) if von == bis else f'{von}–{bis}'
            zvon, zbis = v.get('zeitraumVon', '') or '', v.get('zeitraumBis', '') or ''
            zeit_text = zvon if zvon == zbis else f'{zvon} – {zbis}'
            berichtnr = str(v.get('berichtNr', '') or '').zfill(3) if str(v.get('berichtNr', '')).isdigit() else str(v.get('berichtNr', '') or '')
            zus_text = f"QS-Bericht {berichtnr}: " + (v.get('zusammenfassung', '') or '')
            rows_xml += '<w:tr>' + vcell(nr_text, col_nr, muted=True) + vcell(zeit_text, col_zeit) + vcell(zus_text, col_zus) + '</w:tr>'

        verlauf_tbl_xml = f'''<w:tbl xmlns:w="{W}">
          <w:tblPr>
            <w:tblW w:w="{avail_dxa}" w:type="dxa"/>
            <w:tblLayout w:type="fixed"/>
            <w:tblBorders>
              <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
              <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
              <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            </w:tblBorders>
            <w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:bottom w:w="40" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>
          </w:tblPr>
          <w:tblGrid><w:gridCol w:w="{col_nr}"/><w:gridCol w:w="{col_zeit}"/><w:gridCol w:w="{col_zus}"/></w:tblGrid>
          {rows_xml}
        </w:tbl>'''
        verlauf_tbl = parse_xml(verlauf_tbl_xml)
        verlauf_anchor._p.addprevious(verlauf_tbl)

    # ------------------------------------------------------------------
    # Kopfzeile: Projektname
    # ------------------------------------------------------------------
    projekt_text = bk.get('projekt', '').strip()
    for sec in doc.sections:
        for hdr in (sec.header, sec.first_page_header):
            for p in hdr.paragraphs:
                if 'XY' in p.text:
                    text_runs = [r for r in p.runs if r._r.find(qn('w:drawing')) is None]
                    if not text_runs:
                        continue
                    text_runs[0].text = projekt_text
                    for r in text_runs[1:]:
                        if r.text.strip() and 'Leistungsfeststellung' not in r.text and 'Qualitätssicherung' not in r.text:
                            r.text = ''

    # ------------------------------------------------------------------
    # Deckblatt-Werte
    # ------------------------------------------------------------------
    set_cover_value(doc, 'Verfasser', bk.get('verfasser', ''))
    set_cover_value(doc, 'Datum', bk.get('datum', ''))
    set_cover_value(doc, 'Kunde', bk.get('kunde', ''))
    set_seitenanzahl_field(doc, literal_total=bk.get('combinedTotalPages'))

    # Verteiler
    fill_verteiler(verteiler_table, bk.get('verteiler', ''))

    # Wetterband
    fill_wetterband(doc, bk)

    # Rahmentermine
    fill_rahmentermine(rahmentermine_table, bk)

    # ------------------------------------------------------------------
    # Fußzeile: Dateiname (links) + Seitenzahl (rechts), auf allen Seiten
    # inkl. Deckblatt (Vorlage hat bereits die Abschlusslinie eingebaut).
    # ------------------------------------------------------------------
    dateiname_val = bk.get('dateinameFooter') or bk.get('dateiname') or (
        (bk.get('datum', '').replace('.', '') or 'bericht') + '_' +
        (bk.get('verfasser', '') or 'QS') + '-QS-Bautenstand_' +
        str(bk.get('berichtsNr', '1')).zfill(3) + '.docx')
    dateiname_val = re.sub(r'\.docx$', '', dateiname_val, flags=re.IGNORECASE)
    combined_total_pages = bk.get('combinedTotalPages')
    page_start = bk.get('pageStart')
    if page_start:
        sectPr = doc.sections[0]._sectPr
        old_pgnum = sectPr.find(qn('w:pgNumType'))
        if old_pgnum is not None:
            sectPr.remove(old_pgnum)
        pgnum_el = OxmlElement('w:pgNumType')
        pgnum_el.set(qn('w:start'), str(int(page_start)))
        pgMar = sectPr.find(qn('w:pgMar'))
        if pgMar is not None:
            pgMar.addnext(pgnum_el)
        else:
            sectPr.append(pgnum_el)

    def build_footer_table(ftr):
        sdt = ftr._element.find(f'.//{w("sdt")}')
        if sdt is None:
            return
        sdt_content = sdt.find(qn('w:sdtContent'))
        sdt_paragraphs = sdt_content.findall(qn('w:p')) if sdt_content is not None else []
        rpr_xml = f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="14"/><w:szCs w:val="14"/><w:color w:val="{COL_FOOTER}"/></w:rPr>'
        for rn in (sdt_content.findall(f'.//{w("r")}') if sdt_content is not None else []):
            old_rpr = rn.find(qn('w:rPr'))
            if old_rpr is not None:
                rn.remove(old_rpr)
            rn.insert(0, parse_xml(rpr_xml))
        if sdt_paragraphs:
            target_p = sdt_paragraphs[-1]
            def make_run(inner):
                return parse_xml(f'<w:r xmlns:w="{W}">{rpr_xml}{inner}</w:r>')
            target_p.append(make_run('<w:t xml:space="preserve"> / </w:t>'))
            if combined_total_pages:
                target_p.append(make_run(f'<w:t>{int(combined_total_pages)}</w:t>'))
            else:
                target_p.append(make_run('<w:fldChar w:fldCharType="begin"/>'))
                target_p.append(make_run('<w:instrText>NUMPAGES   \\* MERGEFORMAT</w:instrText>'))
                target_p.append(make_run('<w:fldChar w:fldCharType="separate"/>'))
                target_p.append(make_run('<w:t>1</w:t>'))
                target_p.append(make_run('<w:fldChar w:fldCharType="end"/>'))
            for extra_p in sdt_paragraphs[:-1]:
                sdt_content.remove(extra_p)

        for p_el in list(ftr._element.findall(qn('w:p'))):
            ftr._element.remove(p_el)

        sec = doc.sections[0]
        avail_dxa = sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips
        left_dxa = int(avail_dxa * 0.65)
        right_dxa = avail_dxa - left_dxa

        sdt.getparent().remove(sdt)

        tbl_xml = f'''<w:tbl xmlns:w="{W}">
          <w:tblPr>
            <w:tblW w:w="{avail_dxa}" w:type="dxa"/>
            <w:tblLayout w:type="fixed"/>
            <w:tblBorders>
              <w:top w:val="single" w:sz="6" w:space="4" w:color="{COL_LINE_DARK}"/>
              <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
              <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
              <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            </w:tblBorders>
            <w:tblCellMar><w:top w:w="36" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>
          </w:tblPr>
          <w:tblGrid><w:gridCol w:w="{left_dxa}"/><w:gridCol w:w="{right_dxa}"/></w:tblGrid>
          <w:tr>
            <w:tc><w:tcPr><w:tcW w:w="{left_dxa}" w:type="dxa"/></w:tcPr>
              <w:p><w:pPr><w:jc w:val="left"/></w:pPr>
                <w:r><w:rPr><w:rFonts w:ascii="Barlow" w:hAnsi="Barlow"/><w:sz w:val="14"/><w:szCs w:val="14"/><w:color w:val="{COL_FOOTER}"/></w:rPr>
                <w:t xml:space="preserve">{dateiname_val}</w:t></w:r>
              </w:p>
            </w:tc>
            <w:tc><w:tcPr><w:tcW w:w="{right_dxa}" w:type="dxa"/></w:tcPr></w:tc>
          </w:tr>
        </w:tbl>'''
        tbl = parse_xml(tbl_xml)
        right_tc = tbl.findall(qn('w:tr') + '/' + qn('w:tc'))[1]
        right_tc.append(sdt)
        ftr._element.append(tbl)

    build_footer_table(doc.sections[0].footer)
    build_footer_table(doc.sections[0].first_page_footer)

    # ------------------------------------------------------------------
    # Dokumentationstabelle
    # ------------------------------------------------------------------
    table = dokumentation_table
    tbl = table._tbl
    template_row_tr = copy.deepcopy(table.rows[1]._tr)

    for r in list(table.rows)[1:]:
        tbl.remove(r._tr)

    os.makedirs(tmp_dir, exist_ok=True)
    for i, entry in enumerate(entries, start=start_nr):
        new_tr = copy.deepcopy(template_row_tr)
        tbl.append(new_tr)
        row = table.rows[-1]
        set_nr_cell(row.cells[0], i)
        photo_path = os.path.join(tmp_dir, f"entry_{i}.jpg")
        decode_photo(entry['photoUrl'], photo_path)
        image_height_pt = build_image_cell(row.cells[2], photo_path)
        build_text_cell(row.cells[1], entry, None, image_height_pt)

    if only_content:
        body = doc.element.body
        stop_el = dokumentation_p._p
        to_remove = []
        el = body[0]
        while el is not None and el is not stop_el:
            nxt = el.getnext()
            to_remove.append(el)
            el = nxt
        for el in to_remove:
            body.remove(el)
        doc.sections[0].different_first_page_header_footer = False

    doc.save(out_path)
    settings_el = doc.settings.element
    if settings_el.find(qn('w:updateFields')) is None:
        update_fields_el = OxmlElement('w:updateFields')
        update_fields_el.set(qn('w:val'), 'true')
        settings_el.insert(0, update_fields_el)
    doc.save(out_path)
    return out_path
